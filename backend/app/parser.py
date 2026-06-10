from __future__ import annotations

import re
import subprocess
import tempfile
import sys
from email import policy
from email.parser import BytesParser
from html import unescape
from io import BytesIO
from pathlib import Path

from fastapi import UploadFile


def _strip_html_tags(value: str) -> str:
    no_tags = re.sub(r"<[^>]+>", " ", value)
    clean = re.sub(r"\s+", " ", unescape(no_tags)).strip()
    return clean


def _extract_email_text(data: bytes) -> str:
    message = BytesParser(policy=policy.default).parsebytes(data)
    parts: list[str] = []

    if message.is_multipart():
        for part in message.walk():
            content_type = part.get_content_type()
            if content_type not in {"text/plain", "text/html"}:
                continue
            try:
                payload = part.get_content()
            except LookupError:
                continue
            if not isinstance(payload, str):
                continue
            parts.append(_strip_html_tags(payload) if content_type == "text/html" else payload.strip())
    else:
        payload = message.get_content()
        if isinstance(payload, str):
            content_type = message.get_content_type()
            parts.append(_strip_html_tags(payload) if content_type == "text/html" else payload.strip())

    subject = str(message.get("subject", "")).strip()
    sender = str(message.get("from", "")).strip()
    recipient = str(message.get("to", "")).strip()
    header_block = "\n".join(line for line in [f"Subject: {subject}" if subject else "", f"From: {sender}" if sender else "", f"To: {recipient}" if recipient else ""] if line)
    body = "\n\n".join(part for part in parts if part).strip()
    return "\n\n".join(part for part in [header_block, body] if part).strip()


def _extract_via_textutil(data: bytes, suffix: str) -> str:
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as handle:
        handle.write(data)
        temp_path = Path(handle.name)

    try:
        result = subprocess.run(
            ["textutil", "-convert", "txt", "-stdout", str(temp_path)],
            check=True,
            capture_output=True,
            text=True,
        )
        return result.stdout.strip()
    finally:
        temp_path.unlink(missing_ok=True)


async def extract_text_from_upload(file: UploadFile) -> str:
    data = await file.read()
    filename = (file.filename or "").lower()

    if filename.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="replace")

    if filename.endswith((".html", ".htm")):
        return _strip_html_tags(data.decode("utf-8", errors="replace"))

    if filename.endswith(".eml"):
        return _extract_email_text(data)

    if filename.endswith(".pdf"):
        try:
            import fitz
        except ImportError:
            try:
                from pypdf import PdfReader
            except ImportError as exc:
                raise RuntimeError("PDF parsing requires pypdf. Install backend/requirements.txt.") from exc

            reader = PdfReader(BytesIO(data))
            return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()

        return _extract_pdf_with_pymupdf(data)

    if filename.endswith(".docx"):
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise RuntimeError("DOCX parsing requires python-docx. Install backend/requirements.txt.") from exc

        document = Document(BytesIO(data))
        return "\n".join(paragraph.text for paragraph in document.paragraphs).strip()

def _extract_pdf_with_pymupdf(data: bytes) -> str:
    """Extract PDF using PyMuPDF with block-sorting for better reading order.
    
    Integrated from legacy backend for improved PDF text extraction.
    Sorts blocks by vertical then horizontal coordinate to preserve reading order.
    """
    import fitz
    
    doc = fitz.open(stream=data, filetype="pdf")
    pages_text = []
    
    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        blocks = page.get_text("blocks")
        
        # Sort blocks by y0 (vertical) then x0 (horizontal) to maintain reading order
        try:
            blocks_sorted = sorted(blocks, key=lambda b: (b[1], b[0]))
        except Exception:
            blocks_sorted = blocks
        
        texts = []
        for block in blocks_sorted:
            # Block format: (x0, y0, x1, y1, text, block_no, block_type)
            try:
                text = block[4] if len(block) > 4 else str(block)
            except (IndexError, TypeError):
                text = str(block)
            
            if isinstance(text, str) and text.strip():
                texts.append(text.strip())
        
        page_text = "\n".join(texts)
        if page_text.strip():
            pages_text.append(page_text)
    
    doc.close()
    return "\n\n".join(pages_text).strip()


def extract_text_from_bytes(data: bytes, filename: str | None = None) -> str:
    name = (filename or "").lower()

    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="replace")

    if name.endswith((".html", ".htm")):
        return _strip_html_tags(data.decode("utf-8", errors="replace"))

    if name.endswith(".eml"):
        return _extract_email_text(data)

    if name.endswith(".pdf"):
        try:
            import fitz
        except ImportError:
            try:
                from pypdf import PdfReader
            except ImportError as exc:
                raise RuntimeError("PDF parsing requires pypdf. Install backend/requirements.txt.") from exc

            reader = PdfReader(BytesIO(data))
            return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()

        return _extract_pdf_with_pymupdf(data)

    if name.endswith(".docx"):
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise RuntimeError("DOCX parsing requires python-docx. Install backend/requirements.txt.") from exc

        document = Document(BytesIO(data))
        return "\n".join(paragraph.text for paragraph in document.paragraphs).strip()

    if name.endswith((".doc", ".rtf")):
        if sys.platform != "darwin":
            raise RuntimeError("DOC or RTF parsing requires macOS textutil. Convert to DOCX or PDF.")
        try:
            return _extract_via_textutil(data, suffix=Path(name).suffix)
        except (FileNotFoundError, subprocess.CalledProcessError) as exc:
            raise RuntimeError("DOC or RTF parsing requires macOS textutil.") from exc

    raise ValueError("Unsupported file type. Upload PDF, DOC, DOCX, EML, HTML, Markdown, RTF, or TXT.")

    if filename.endswith((".doc", ".rtf")):
        if sys.platform != "darwin":
            raise RuntimeError("DOC or RTF parsing requires macOS textutil. Convert to DOCX or PDF.")
        try:
            return _extract_via_textutil(data, suffix=Path(filename).suffix)
        except (FileNotFoundError, subprocess.CalledProcessError) as exc:
            raise RuntimeError("DOC or RTF parsing requires macOS textutil.") from exc

    raise ValueError("Unsupported file type. Upload PDF, DOC, DOCX, EML, HTML, Markdown, RTF, or TXT.")
